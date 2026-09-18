"""Deterministic documents: model text is data, never HTML, code or a template."""
import sys, json, os, hashlib, zipfile, re
from pathlib import Path
from io import BytesIO
from html import escape
from datetime import date
from PIL import Image, ImageOps
import pypdfium2 as pdfium
from pypdf import PdfReader, PdfWriter
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.utils import ImageReader
from fontTools.ttLib import TTFont as FontFile
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn

Image.MAX_IMAGE_PIXELS = 30_000_000
ROOT = Path(__file__).resolve().parent

def pages(filename, limit=12):
    source = Path(filename)
    if source.read_bytes()[:5] == b'%PDF-':
        pdf = pdfium.PdfDocument(str(source))
        if len(pdf) > limit or len(pdf) == 0:
            raise ValueError(f'PDF 须为 1–{limit} 页，不能截断识别')
        try:
            for n in range(len(pdf)):
                page = pdf[n]
                size = page.get_size()
                scale = min(1.7, 2400 / max(size))
                bitmap = page.render(scale=scale)
                image = bitmap.to_pil().convert('RGB').copy()
                bitmap.close(); page.close()
                yield image
        finally:
            pdf.close()
    else:
        with Image.open(source) as raw:
            if getattr(raw, 'n_frames', 1) != 1:
                raise ValueError('不支持多帧图片，请转换为 PDF 或单张 PNG')
            image = ImageOps.exif_transpose(raw).convert('RGB')
            image.thumbnail((2400, 2400))
            yield image.copy()

def prepare_images(data, out):
    result = []
    for i, image in enumerate(pages(data['path'], 8)):
        filename = out / f'page-{i+1}.png'
        image.save(filename)
        result.append(str(filename))
    return {'images': result}

def make_packet(data, out):
    font = ROOT / 'assets' / 'fonts' / 'NotoSansSC-Regular.ttf'
    pdfmetrics.registerFont(TTFont('CN', str(font)))
    cmap = FontFile(font).getBestCmap()
    record = data['record']; purpose = data['purpose'].strip()
    if not purpose or len(purpose) > 6000:
        raise ValueError('用途说明须为 1–6000 字')
    title = 'ChatGPT 订阅费用报销说明'
    fx = record.get('exchangeRate') or {}
    amount = record['amount'] + ' ' + record['currency']
    fields = [('账单月份', record['billingMonth']), ('发票编号', record['invoiceNumber']),
              ('发票日期', record['date']), ('原币金额', amount), ('申报人民币', 'CNY ' + record['claimedCNY'])]
    if record['currency'] != 'CNY':
        fields.extend([('汇率日期', fx['date']), ('汇率口径', '中国银行 · 中行折算价'),
                       ('换算过程', f"{record['amount']} × {fx['quotedRate']} ÷ 100 = CNY {record['claimedCNY']}")])
    basis = '汇率按发票日期取中行折算价（每100外币兑人民币），沿用用户确认的同类已获批报销先例。'
    if record['currency'] == 'CNY': basis = '人民币费用按原币金额申报，无需外汇换算。'
    all_text = title + purpose + basis + ''.join(k+v for k,v in fields)
    if any(ord(c) not in cmap for c in all_text if not c.isspace()):
        raise ValueError('说明包含中文字体不支持的字符，请删除表情或特殊符号后重试')
    normal = ParagraphStyle('body', fontName='CN', fontSize=11, leading=19, wordWrap='CJK', spaceAfter=8)
    heading = ParagraphStyle('heading', parent=normal, fontSize=13, leading=22, spaceBefore=14, spaceAfter=8)
    doc_title = ParagraphStyle('title', parent=normal, fontSize=19, leading=27, spaceAfter=22)
    small = ParagraphStyle('small', parent=normal, fontSize=9, leading=15)
    paragraph = lambda value, style=normal: Paragraph(escape(str(value)).replace('\n','<br/>'), style)
    rows = [[paragraph(k, small), paragraph(v)] for k,v in fields]
    table = Table(rows, colWidths=[100, 385], hAlign='LEFT')
    table.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),0),('RIGHTPADDING',(0,0),(-1,-1),10),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
    flow = [paragraph(title, doc_title), table, paragraph('科研或工作用途', heading)]
    flow += [paragraph(p) for p in purpose.split('\n') if p.strip()]
    flow += [paragraph('金额与附件说明', heading), paragraph(basis), paragraph('后附发票、付款凭证及汇率查询截图。原件另行保存在报销台账，附件按原件内容编排。', small)]
    cover = out / 'statement.pdf'
    SimpleDocTemplate(str(cover), pagesize=A4, leftMargin=55, rightMargin=55, topMargin=48, bottomMargin=48, title=title, author='').build(flow)

    docx = Document()
    sec = docx.sections[0]; sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=sec.bottom_margin=Cm(1.8);sec.left_margin=sec.right_margin=Cm(2)
    for style_name in ['Normal','Title','Heading 1']:
        style=docx.styles[style_name];style.font.name='Noto Sans SC';style.font.size=Pt(11 if style_name=='Normal' else 18 if style_name=='Title' else 13)
        style.font.color.rgb=RGBColor(0,0,0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans SC')
    docx.add_paragraph(title,'Title')
    for k,v in fields: docx.add_paragraph(k+'：'+v)
    docx.add_heading('科研或工作用途',1)
    for p in purpose.split('\n'):
        if p.strip(): docx.add_paragraph(p)
    docx.add_heading('金额与附件说明',1);docx.add_paragraph(basis)
    docx.add_paragraph('发票、付款凭证及汇率截图详见同批整合 PDF。')
    docx.core_properties.author='';docx.core_properties.title=title
    word = out / 'statement.docx';docx.save(word)

    writer = PdfWriter(); writer.append(str(cover)); mappings=[]
    labels={'invoice':'原始发票','payment':'付款凭证','exchangeRate':'汇率查询截图','purposeEvidence':'用途补充材料'}
    for source_index, source in enumerate(data['sources']):
        content=Path(source['path']).read_bytes()
        if hashlib.sha256(content).hexdigest()!=source['sha256']: raise ValueError('生成前原件哈希校验失败')
        first=len(writer.pages)+1
        for page_index, image in enumerate(pages(source['path'])):
            dimensions=landscape(A4) if image.width/image.height>1.4 else A4
            page_file=out/f'attachment-{source_index}-{page_index}.pdf'
            canvas=Canvas(str(page_file),pagesize=dimensions)
            width,height=dimensions;canvas.setFont('CN',11)
            canvas.drawString(32,height-30,labels.get(source['role'],'附件')+f' · {source_index+1}.{page_index+1}')
            available_w,available_h=width-64,height-90
            scale=min(available_w/image.width,available_h/image.height)
            iw,ih=image.width*scale,image.height*scale
            canvas.drawImage(ImageReader(image),(width-iw)/2,40+(available_h-ih)/2,width=iw,height=ih)
            canvas.save();writer.append(str(page_file))
        mappings.append({'id':source['id'],'role':source['role'],'sha256':source['sha256'],'firstPage':first,'lastPage':len(writer.pages)})
    combined=out/'application.pdf'
    writer.add_metadata({'/Title':title,'/Author':'','/Subject':record['invoiceNumber']})
    with combined.open('wb') as stream:writer.write(stream)
    reader=PdfReader(combined)
    cover_text='\n'.join(p.extract_text() or '' for p in reader.pages[:len(PdfReader(cover).pages)])
    if record['invoiceNumber'] not in cover_text or '科研或工作用途' not in cover_text: raise ValueError('PDF 中文文字校验失败')
    previews=[]
    pdf=pdfium.PdfDocument(str(combined))
    for n in range(len(pdf)):
        page=pdf[n];bitmap=page.render(scale=1.2);preview=out/f'preview-{n+1}.png';bitmap.to_pil().save(preview);bitmap.close();page.close();previews.append(str(preview))
    pdf.close()
    fonts=reader.pages[0]['/Resources']['/Font'].get_object()
    embedded=any('/FontFile2' in font.get_object()['/FontDescriptor'].get_object() for font in fonts.values() if '/FontDescriptor' in font.get_object())
    if not embedded: raise ValueError('PDF 中文字体未嵌入')
    return {'pdf':str(combined),'docx':str(word),'previews':previews,'pages':len(reader.pages),'fontEmbedded':True,'sourcePageMap':mappings,'templateVersion':'payment-evidence-v1'}

def make_zip(data,out):
    target=out/'reimbursement-package.zip'
    if not 1<=len(data['files'])<=30:raise ValueError('ZIP 需包含 1–30 份 PDF')
    with zipfile.ZipFile(target,'w',compression=zipfile.ZIP_DEFLATED) as archive:
        for index,item in enumerate(data['files']):
            safe=re.sub(r'[^A-Za-z0-9._-]','_',item['name'])
            filename=f'{index+1:02d}-{safe}'
            if not filename.lower().endswith('.pdf'):raise ValueError('申报 ZIP 只包含 PDF')
            content=Path(item['path']).read_bytes()
            if hashlib.sha256(content).hexdigest()!=item['sha256']:raise ValueError('PDF 哈希与登记不一致')
            archive.writestr(filename,content)
    with zipfile.ZipFile(target) as archive:
        if archive.testzip():raise ValueError('ZIP 完整性校验失败')
    return {'zip':str(target),'files':len(data['files'])}

if __name__=='__main__':
    command,request,result=sys.argv[1:]
    request=Path(request);out=request.parent
    data=json.loads(request.read_text(encoding='utf-8'))
    action={'images':prepare_images,'packet':make_packet,'zip':make_zip}[command]
    Path(result).write_text(json.dumps(action(data,out),ensure_ascii=False),encoding='utf-8')
